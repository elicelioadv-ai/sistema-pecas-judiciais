"""
motor.py - o "cérebro" do programa.

Responsabilidades:
  1. Carregar/salvar a base legal e os modelos de peça.
  2. Chamar a IA (Gemini ou Claude) para organizar fatos e redigir a peça.
  3. CONFERIR o resultado: citações, datas e valores (anti-invenção).
  4. Apoiar histórico de casos, transcrição de áudio e exportação (Word/PDF).
"""
from __future__ import annotations

import hashlib
import hmac
import json
import logging
import os
import re
import time
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from dotenv import load_dotenv

PASTA = Path(__file__).parent
load_dotenv(PASTA / ".env")

PASTA_BASE = PASTA / "base"
PASTA_MODELOS = PASTA / "modelos"
PASTA_SAIDA = PASTA / "saida"
PASTA_HISTORICO = PASTA / "historico"
PASTA_RASCUNHOS = PASTA / "rascunhos"
for _p in (PASTA_BASE, PASTA_MODELOS, PASTA_SAIDA, PASTA_HISTORICO, PASTA_RASCUNHOS):
    _p.mkdir(exist_ok=True)

_LOG = logging.getLogger("gerador")
if not _LOG.handlers:
    _h = logging.FileHandler(PASTA / "gerador.log", encoding="utf-8")
    _h.setFormatter(logging.Formatter("%(asctime)s [%(levelname)s] %(message)s"))
    _LOG.addHandler(_h)
    _LOG.setLevel(logging.INFO)

STATUS_VALIDOS = ("verificado", "conferir", "vazio")


def _backup_drive(caminho, subpasta="pecas"):
    """Copia o arquivo salvo para a nuvem (GitHub ou Drive, se configurado). Nunca quebra o app."""
    try:
        import nuvem

        if nuvem.disponivel():
            nuvem.enviar(caminho, subpasta)
    except Exception:
        _LOG.warning("Backup na nuvem falhou (%s)", caminho, exc_info=True)


def _backup_remover(nome_arquivo, subpasta="casos"):
    """Remove da nuvem o arquivo apagado aqui (mantém o backup fiel). Nunca quebra o app."""
    try:
        import nuvem

        if nuvem.disponivel():
            nuvem.remover(nome_arquivo, subpasta)
    except Exception:
        _LOG.warning("Remoção na nuvem falhou (%s)", nome_arquivo, exc_info=True)


# ---------------------------------------------------------------- configuração
PROVEDOR = os.getenv("PROVEDOR", "gemini").strip().lower()
_MODELOS_PADRAO = {"gemini": "gemini-2.5-flash", "anthropic": "claude-sonnet-4-5"}
MODELO = os.getenv("MODELO", "").strip() or _MODELOS_PADRAO.get(PROVEDOR, "")


def _recarregar_config():
    global PROVEDOR, MODELO
    PROVEDOR = os.getenv("PROVEDOR", "gemini").strip().lower()
    MODELO = os.getenv("MODELO", "").strip() or _MODELOS_PADRAO.get(PROVEDOR, "")


def _limpar_chave(nome):
    return os.getenv(nome, "").strip().strip('"').strip("'")


def modo_uso():
    """Situação atual do provedor e do tratamento de dados (para avisos na tela)."""
    return {
        "provedor": PROVEDOR,
        "modelo": MODELO,
        "modo_chave": os.getenv("MODO_CHAVE", "gemini").strip().lower(),
        "regiao": os.getenv("REGIAO", "").strip(),
        "uso_pago": os.getenv("USO_PAGO", "nao").strip().lower() in ("sim", "yes", "true", "1"),
    }


def atualizar_env(valores):
    """Grava/atualiza chaves no .env preservando as demais (inclusive a chave da API)."""
    caminho = PASTA / ".env"
    linhas = caminho.read_text(encoding="utf-8").splitlines() if caminho.exists() else []
    restantes = dict(valores)
    novas = []
    for linha in linhas:
        m = re.match(r"\s*([A-Za-z_][A-Za-z0-9_]*)\s*=", linha)
        if m and m.group(1) in restantes:
            novas.append(f"{m.group(1)}={restantes.pop(m.group(1))}")
        else:
            novas.append(linha)
    novas.extend(f"{k}={v}" for k, v in restantes.items())
    caminho.write_text("\n".join(novas) + "\n", encoding="utf-8")
    load_dotenv(caminho, override=True)
    _recarregar_config()
    _LOG.info("Configuração atualizada no .env: %s", ", ".join(valores))
    return caminho


# ---------------------------------------------------------------- acesso (senha)
def senha_acesso():
    """Senha armazenada no .env (pode ser 'sha256:<hex>' ou texto puro)."""
    return _limpar_chave("SENHA_ACESSO")


def senha_ativa():
    return bool(senha_acesso())


def usuario_acesso():
    """Usuário configurado no .env (ACESSO_USUARIO). Vazio = ainda não definido."""
    return _limpar_chave("ACESSO_USUARIO")


def _hash_senha(senha):
    import hashlib

    return "sha256:" + hashlib.sha256(senha.encode("utf-8")).hexdigest()


def verificar_senha(senha):
    """Compara a senha digitada com a armazenada (sem vazar a armazenada)."""
    armazenada = senha_acesso()
    if not armazenada:
        return False
    import hmac

    if armazenada.startswith("sha256:"):
        return hmac.compare_digest(_hash_senha((senha or "").strip()), armazenada)
    return hmac.compare_digest((senha or "").strip(), armazenada)


def verificar_acesso(usuario, senha):
    """Confere usuário e senha juntos no portão de entrada."""
    import hmac

    usuario_ok = hmac.compare_digest((usuario or "").strip().casefold(), usuario_acesso().strip().casefold())
    return usuario_ok and verificar_senha(senha)


def definir_acesso(usuario, senha):
    """Define usuário e senha de acesso no .env (senha guardada com hash)."""
    usuario = (usuario or "").strip()
    senha = (senha or "").strip()
    if len(usuario) < 2:
        raise ValueError("Informe um usuário (mínimo 2 caracteres).")
    if len(senha) < 4:
        raise ValueError("A senha precisa ter ao menos 4 caracteres.")
    return atualizar_env({"ACESSO_USUARIO": usuario, "SENHA_ACESSO": _hash_senha(senha)})


def definir_senha(senha):
    """Define e grava a senha de acesso no .env (guardada com hash)."""
    senha = (senha or "").strip()
    if len(senha) < 4:
        raise ValueError("A senha precisa ter ao menos 4 caracteres.")
    return atualizar_env({"SENHA_ACESSO": _hash_senha(senha)})


# ---------------------------------------------------------------- anonimização (LGPD)
class Anonimizador:
    """Troca dados pessoais por marcadores ([PESSOA_1], [CPF_1]...) antes de enviar à IA.

    O texto devolvido pela IA é restaurado com os dados originais, de modo que o
    advogado sempre vê a versão real. É uma camada de proteção, e não substitui a
    revisão humana nem a avaliação jurídica de conformidade (LGPD).
    """

    _RX = [
        ("CNPJ", re.compile(r"\b\d{2}\.?\d{3}\.?\d{3}/?\d{4}-?\d{2}\b")),
        ("CPF", re.compile(r"\b\d{3}\.?\d{3}\.?\d{3}-?\d{2}\b")),
        ("CEP", re.compile(r"\b\d{5}-\d{3}\b")),
        ("EMAIL", re.compile(r"\b[\w.+-]+@[\w-]+\.[\w.-]+\b")),
        ("TELEFONE", re.compile(r"\b\d{4,5}[-.]\d{4}\b")),
    ]
    _RX_ROTULO = re.compile(
        r"(?i)\b(autor(?:a)?|réu|ré|requerente|requerido(?:a)?|cliente|testemunha|acompanhante|"
        r"apelante|apelado(?:a)?|executado(?:a)?|exequente|nome|respons[áa]vel)\b\s*[:\-–]\s*"
        r"([a-záàâãéèêíìóòôõúùçñ][\w’'.\-]*(?:\s+(?:d[aeo]s?\s+)?[\w’'.\-]+){0,4})"
    )
    _DESCRITORES = ("brasileir", "solteir", "casad", "divorciad", "viuv", "professor", "portador",
                    "residente", "domiciliad", "inscrit", "autônom", "aposentad", "estudant",
                    "maior", "menor", "nascid", "filh", "comerciant", "empresár", "servidor")
    _ORGANIZACAO = re.compile(
        r"(?i)\b(s\.?\s?a\.?|ltda|s\s?/?\s?a|eireli|me|epp|cia\.?|companhia|banco|seguradora|"
        r"empresa|aerolinhas|airlines|transportadora)\b"
    )
    _RX_MARCADOR = re.compile(r"\[\s*(?:PESSOA|ORGANIZACAO|CPF|CNPJ|EMAIL|CEP|TELEFONE)[\s_#]*\d+\s*\]", re.I)

    def __init__(self, termos_extras=None):
        self.termos_extras = [t.strip() for t in (termos_extras or []) if t and t.strip()]

    def _categoria(self, nome):
        return "ORGANIZACAO" if self._ORGANIZACAO.search(nome) else "PESSOA"

    def _limpar_nome(self, nome):
        partes = []
        for palavra in nome.strip(" .,;").split():
            if palavra.lower().strip(".,;").startswith(self._DESCRITORES):
                break
            partes.append(palavra)
        return " ".join(partes).strip(" .,;")

    def mascarar(self, texto):
        """Devolve (texto_mascarado, mapa) onde mapa liga o marcador ao valor original."""
        mapa, contador, resultado = {}, defaultdict(int), texto or ""
        entidades = []
        for m in self._RX_ROTULO.finditer(resultado):
            nome = self._limpar_nome(m.group(2))
            if len(nome) >= 3:
                entidades.append((nome, self._categoria(nome)))
        entidades.extend((t, self._categoria(t)) for t in self.termos_extras)

        for valor, categoria in sorted(entidades, key=lambda x: len(x[0]), reverse=True):
            for achado in sorted(set(re.findall(re.escape(valor), resultado, flags=re.IGNORECASE)), key=len, reverse=True):
                contador[categoria] += 1
                ph = f"[{categoria}_{contador[categoria]}]"
                mapa[ph] = achado
                resultado = re.sub(re.escape(achado), ph, resultado)

        for categoria, padrao in self._RX:
            for achado in sorted(set(padrao.findall(resultado)), key=len, reverse=True):
                contador[categoria] += 1
                ph = f"[{categoria}_{contador[categoria]}]"
                mapa[ph] = achado
                resultado = re.sub(re.escape(achado), ph, resultado)

        return resultado, mapa

    def restaurar(self, texto, mapa):
        if not mapa:
            return texto

        def rep(m):
            return mapa.get(f"[{m.group(1).upper()}_{m.group(2)}]", m.group(0))

        return re.sub(r"\[\s*([A-Za-zÀ-ú]+)[\s_#]*(\d+)\s*\]", rep, texto or "")

    def marcadores_restantes(self, texto):
        return self._RX_MARCADOR.findall(texto or "")


# ---------------------------------------------------------------- base legal
def listar_bases():
    return sorted(p.stem for p in PASTA_BASE.glob("*.json"))


def carregar_base(nome):
    with open(PASTA_BASE / f"{nome}.json", encoding="utf-8") as f:
        return json.load(f)


def salvar_base(nome, itens):
    if not re.fullmatch(r"[A-Za-z0-9_\-]+", nome or ""):
        raise ValueError("O nome da base deve conter apenas letras, números, _ ou -.")
    caminho = PASTA_BASE / f"{nome}.json"
    caminho.write_text(json.dumps(itens, ensure_ascii=False, indent=2), encoding="utf-8")
    _LOG.info("Base salva: %s (%d itens)", nome, len(itens))
    _backup_drive(caminho, "bases")
    return caminho


def apagar_base(nome):
    if not re.fullmatch(r"[A-Za-z0-9_\-]+", nome or ""):
        raise ValueError("O nome da base deve conter apenas letras, números, _ ou -.")
    caminho = PASTA_BASE / f"{nome}.json"
    if caminho.exists():
        caminho.unlink()
    _backup_remover(f"{nome}.json", "bases")
    _LOG.info("Base apagada: %s", nome)
    return nome


def novo_item_base(id_, referencia, texto, status="conferir", fonte=""):
    if status not in STATUS_VALIDOS:
        raise ValueError(f"status inválido: {status!r}. Use um de {STATUS_VALIDOS}.")
    return {"id": id_, "referencia": referencia, "status": status, "texto": texto, "fonte": fonte}


# ---------------------------------------------------------------- modelos de peça
def listar_modelos():
    return sorted(p.stem for p in PASTA_MODELOS.glob("*.md"))


def carregar_modelo(nome):
    return (PASTA_MODELOS / f"{nome}.md").read_text(encoding="utf-8")


def identificar_modelo(fatos, modelos):
    """Sugere o modelo de peça mais adequado aos fatos (heurística local).

    Retorna o nome de um modelo da lista ou o primeiro modelo como padrão.
    """
    if not modelos:
        return None
    texto = (fatos or "").lower()
    regras = [
        ("recurso_inominado", ["recurso inominado", "recurso inominado do jec",
                               "insurgir contra a sentença", "recorrer da sentença", "recurso cível"]),
        ("contestacao", ["contestação", "contestar", "defesa", "tempestividade", "réplica", "impugnação"]),
        ("apelacao", ["apelação", "apelar da sentença", "reexame necessário", "sentença impugnar"]),
        ("embargos", ["embargos de declaração", "embargo de declaração", "obscuridade",
                      "contradição", "omissão", "ambiguidade"]),
    ]
    for slug, marcadores in regras:
        if any(m in texto for m in marcadores):
            for mod in modelos:
                if slug in mod.lower():
                    return mod
            return modelos[0]
    return modelos[0]


# ---------------------------------------------------------------- IA
def _cliente_gemini():
    from google import genai

    chave = _limpar_chave("GEMINI_API_KEY")
    if not chave or "cole_" in chave:
        raise RuntimeError("Chave do Gemini não configurada. Veja o arquivo .env")
    # tempo maximo de 2 minutos por chamada e sem retentativa automatica da
    # biblioteca (o retry fica sob nosso controle, para a tela nao travar)
    opcoes = {"timeout": 120000}
    try:
        from google.genai import types as _types
        opcoes = {"timeout": 120000,
                  "retry_options": _types.HttpRetryOptions(attempts=1, initial_delay=0.0,
                                                          max_delay=0.0, exp_base=1.0, jitter=0.0)}
    except Exception:
        pass
    if os.getenv("MODO_CHAVE", "gemini").strip().lower() == "vertex":
        regiao = os.getenv("REGIAO", "").strip()
        if regiao:
            try:
                return genai.Client(vertexai=True, api_key=chave, location=regiao, http_options=opcoes)
            except Exception:
                _LOG.warning("Falha ao usar a região '%s'; tentando sem região.", regiao)
        return genai.Client(vertexai=True, api_key=chave, http_options=opcoes)
    return genai.Client(api_key=chave, http_options=opcoes)


def _segredo_sessao():
    """Chave do cookie de sessao (fica so no servidor; muda se a senha mudar)."""
    base = senha_acesso() or usuario_acesso() or "pecas-local"
    return hashlib.sha256(("pecas-sessao-v1|" + base).encode("utf-8")).hexdigest()


def token_sessao(usuario, admin, dias=30):
    """Cria o cookie de sessao (assinado): mantem o acesso apos uma atualizacao."""
    exp = int(time.time()) + int(dias) * 24 * 3600
    dados = f"{usuario}|{'1' if admin else '0'}|{exp}"
    assinatura = hmac.new(_segredo_sessao().encode("utf-8"), dados.encode("utf-8"),
                           hashlib.sha256).hexdigest()
    return f"{dados}|{assinatura}"


def ler_token_sessao(token):
    """Devolve {usuario, admin} se o cookie for valido; senao None."""
    try:
        usuario, admin, exp, assinatura = (token or "").split("|")
        if int(exp) < time.time():
            return None
        dados = f"{usuario}|{admin}|{exp}"
        esperado = hmac.new(_segredo_sessao().encode("utf-8"), dados.encode("utf-8"),
                            hashlib.sha256).hexdigest()
        if not hmac.compare_digest(assinatura, esperado):
            return None
        return {"usuario": usuario, "admin": admin == "1"}
    except Exception:
        return None


def mensagem_amigavel(erro):
    msg = str(erro)
    baixo = msg.lower()
    if "503" in msg or "unavailable" in baixo or "high demand" in baixo or "sobrecarregada" in baixo:
        return ("A IA está sobrecarregada agora (o Google está com muita demanda). "
                "Nada foi perdido: o que você escreveu continua salvo. "
                "Tente de novo em alguns minutos.")
    if "429" in msg or "resource_exhausted" in baixo or "cota" in baixo:
        return ("A cota gratuita da IA acabou por hoje. Aguarde o limite se resetar ou "
                "ative faturamento/Vertex AI nas ⚙️ Configurações.")
    if "timeout" in baixo or "timed out" in baixo or "deadline" in baixo:
        return ("A IA demorou demais para responder e foi interrompida. Nada foi perdido. "
                "Tente de novo em alguns minutos.")
    if "api key" in baixo or "key not valid" in baixo or "api_key" in baixo:
        return "A chave da IA não está válida. Confira nas ⚙️ Configurações (ou nos Secrets da nuvem)."
    if "connection" in baixo or "network" in baixo or "ssl" in baixo:
        return "A conexão com a IA falhou. Verifique a internet e tente de novo."
    return msg


def _chamar_ia(system, contents, max_tokens=6000):
    if PROVEDOR == "anthropic":
        return _chamar_claude(system, contents, max_tokens)
    return _chamar_gemini(system, contents, max_tokens)


def _chamar_gemini(system, contents, max_tokens):
    import time

    from google.genai import types

    cliente = _cliente_gemini()
    config = types.GenerateContentConfig(
        system_instruction=system,
        temperature=0.2,
        max_output_tokens=max_tokens * 3,  # margem: alguns modelos gastam tokens "pensando"
    )
    for tentativa in range(3):
        try:
            resp = cliente.models.generate_content(model=MODELO, contents=contents, config=config)
            if not resp.text:
                raise RuntimeError("A IA não devolveu texto (resposta bloqueada ou vazia). Tente de novo.")
            return resp.text
        except Exception as e:
            msg = str(e)
            _LOG.warning("Falha na chamada à IA (tentativa %d): %s", tentativa + 1, msg)
            if "PerDay" in msg or "free_tier_requests" in msg:
                raise RuntimeError(
                    "Cota gratuita da IA esgotada (limite de requisições POR DIA). "
                    "Aguarde o prazo indicado pelo Google, troque o MODELO no .env "
                    "ou ative faturamento/Vertex AI nas ⚙️ Configurações."
                ) from e
            # servidor sobrecarregado (o Google está com alta demanda): tenta uma
            # vez e, se continuar, avisa de forma clara (antes a tela ficava travada)
            if any(x in msg for x in ("503", "UNAVAILABLE", "high demand", "overloaded")):
                if tentativa < 1:
                    time.sleep(10)
                    continue
                raise RuntimeError(
                    "A IA está sobrecarregada agora (o Google está com muita demanda). "
                    "Nada foi perdido: o que você escreveu continua salvo. "
                    "Tente de novo em alguns minutos."
                ) from e
            if "429" in msg and tentativa < 2:
                time.sleep(20)
                continue
            raise


def _chamar_claude(system, contents, max_tokens):
    import anthropic

    if not MODELO:
        raise RuntimeError("Defina MODELO no .env (ex.: claude-sonnet-4-5) para usar o provedor anthropic.")
    chave = _limpar_chave("ANTHROPIC_API_KEY")
    if not chave or "cole_sua_chave" in chave:
        raise RuntimeError("Chave da Anthropic não configurada. Veja o arquivo .env")
    if not isinstance(contents, str):
        raise RuntimeError("Este provedor só aceita texto; a transcrição de voz exige o Gemini.")
    cliente = anthropic.Anthropic(api_key=chave)
    resp = cliente.messages.create(model=MODELO, max_tokens=max_tokens, system=system,
                                   messages=[{"role": "user", "content": contents}])
    return "".join(b.text for b in resp.content if b.type == "text")


REGRAS_ESCRITA = """Você é revisor de português do Brasil. Corrija APENAS erros de ortografia, acentuação e gramática básica.
REGRAS ABSOLUTAS:
- NÃO mude o sentido, NÃO reescreva, NÃO troque por sinônimos, NÃO mexa em nomes próprios, datas, valores, números, termos jurídicos ou nas marcações de formatação (**, __, ==, ~~, ##, - ).
- Só acerte o que estiver realmente errado. Se o texto estiver correto, não invente problema.
Responda EXATAMENTE com um JSON e nada mais:
{"correcoes":[{"errado":"trecho errado como está no texto","certo":"trecho correto","motivo":"explicação curtinha"}]}
Se não houver erro algum, responda: {"correcoes":[]}
"""


def aplicar_correcao(texto, errado, certo):
    """Troca a primeira ocorrencia da palavra errada pela correta. Devolve None se nao houver."""
    texto = texto or ""
    if not errado or not certo or errado not in texto:
        return None
    return texto.replace(errado, certo, 1)


def _extrair_correcoes(saida):
    """Pega a lista de correções que a IA devolveu (aceita JSON puro ou com cercas)."""
    texto = (saida or "").strip()
    if texto.startswith("```"):
        texto = re.sub(r"^```[a-zA-Z]*\s*", "", texto)
        texto = re.sub(r"```\s*$", "", texto).strip()
    itens = []
    try:
        obj = json.loads(texto)
        itens = obj.get("correcoes", []) if isinstance(obj, dict) else []
    except Exception:
        bloco = re.search(r"\{.*\}", texto, re.S)
        if bloco:
            try:
                obj = json.loads(bloco.group(0))
                itens = obj.get("correcoes", []) if isinstance(obj, dict) else []
            except Exception:
                itens = []
    limpas = []
    for item in itens:
        if not isinstance(item, dict):
            continue
        errado = str(item.get("errado", "")).strip()
        certo = str(item.get("certo", "")).strip()
        motivo = str(item.get("motivo", "")).strip()
        if errado and certo and errado != certo:
            limpas.append({"errado": errado, "certo": certo, "motivo": motivo})
    return limpas


def revisar_escrita(texto, anonimizador=None):
    """Procura erros de escrita e devolve [{errado, certo, motivo}]."""
    entrada = _texto_para_ia(texto)
    mapa = {}
    if anonimizador is not None:
        entrada, mapa = anonimizador.mascarar(entrada)
    saida = _chamar_ia(REGRAS_ESCRITA, entrada, 1000)
    if anonimizador is not None:
        saida = anonimizador.restaurar(saida, mapa)
    return _extrair_correcoes(saida)


REGRAS_FATOS = """Você é assistente de um escritório de advocacia brasileiro.
Tarefa: organizar a narrativa bruta de um cliente em FATOS claros, em ordem cronológica, linguagem jurídica simples.
REGRAS ABSOLUTAS:
- NÃO invente nem acrescente nenhum fato, data, hora, valor ou nome que não esteja no texto recebido.
- NÃO faça argumentação jurídica nem cite leis.
- Se algo importante estiver faltando ou ambíguo, NÃO preencha: liste em PENDÊNCIAS como pergunta ao cliente.
Responda EXATAMENTE neste formato:
=== FATOS ===
(texto dos fatos)
=== PENDENCIAS ===
(uma pergunta por linha, começando com '- ')"""


def limpar_fatos(dados_caso, narrativa, anonimizador=None):
    entrada = (f"DADOS DO CASO:\n{dados_caso}\n\nNARRATIVA BRUTA DO CLIENTE:\n"
               f"{_texto_para_ia(narrativa)}")
    mapa = {}
    if anonimizador is not None:
        entrada, mapa = anonimizador.mascarar(entrada)
    saida = _chamar_ia(REGRAS_FATOS, entrada, 3000)
    if anonimizador is not None:
        saida = anonimizador.restaurar(saida, mapa)
    partes = saida.split("=== PENDENCIAS ===")
    fatos = partes[0].replace("=== FATOS ===", "").strip()
    pendencias = partes[1].strip() if len(partes) > 1 else ""
    return fatos, pendencias


REGRAS_ATUALIZAR_FATOS = """Você é assistente de um escritório de advocacia brasileiro.
Tarefa: incorporar as RESPOSTAS do cliente aos FATOS já organizados, atualizando-os.
REGRAS ABSOLUTAS:
- Use SOMENTE as informações recebidas (fatos, pendências e respostas). NÃO invente data, valor, nome ou circunstância.
- NÃO faça argumentação jurídica nem cite leis.
- Liste em PENDÊNCIAS apenas as perguntas que continuarem SEM resposta.
Responda EXATAMENTE neste formato:
=== FATOS ===
(texto atualizado dos fatos)
=== PENDENCIAS ===
(uma pergunta por linha, começando com '- '; se não houver, escreva 'Nada pendente.')"""


def incorporar_respostas(dados_caso, fatos, pendencias, respostas, anonimizador=None):
    entrada = (
        f"DADOS DO CASO:\n{dados_caso}\n\nFATOS ATUAIS:\n{fatos}\n\n"
        f"PENDÊNCIAS (perguntas):\n{pendencias}\n\nRESPOSTAS DO CLIENTE:\n{respostas}"
    )
    mapa = {}
    if anonimizador is not None:
        entrada, mapa = anonimizador.mascarar(entrada)
    saida = _chamar_ia(REGRAS_ATUALIZAR_FATOS, entrada, 3500)
    if anonimizador is not None:
        saida = anonimizador.restaurar(saida, mapa)
    partes = saida.split("=== PENDENCIAS ===")
    novos_fatos = partes[0].replace("=== FATOS ===", "").strip()
    pendencias = partes[1].strip() if len(partes) > 1 else ""
    if pendencias.lower().startswith("nada pendente"):
        pendencias = ""
    return novos_fatos, pendencias


REGRAS_PECA = """Você redige peças para um advogado brasileiro, que revisará tudo antes de protocolar.
REGRAS ABSOLUTAS (violá-las é falha grave):
1. FATOS: use SOMENTE os fatos fornecidos. Nunca invente data, valor, nome, documento ou circunstância. Dado ausente => escreva [PREENCHER: o que falta]. PARTES: reproduza em toda a peça os termos escolhidos pelo advogado nos DADOS DO CASO (ex.: requerente/requerido; autor/réu em ação criminal; recorrente/recorrido em recurso) adaptando o modelo; se a parte requerente/requerida não constar, identifique-a nos FATOS (quem pede e contra quem) e use-a no cabeçalho; se não houver como identificar, escreva [PREENCHER: qualificação da parte].
2. FUNDAMENTOS — ORDEM LÓGICA: em cada tese, cite as autoridades nesta sequência, usando SOMENTE itens que existam na BASE LEGAL e marcando cada citação com [FONTE:id] no ponto exato:
   (i) o dispositivo legal (lei/artigo);
   (ii) a jurisprudência do tribunal local (TJ);
   (iii) o STJ (precedente/recurso) e/ou as súmulas;
   (iv) normativos do CNJ e resoluções de agências reguladoras (ex.: ANAC);
   (v) outras fontes cabíveis.
   Escreva a lei e o artigo UMA única vez (ex.: "art. 14 do CDC"); NÃO repita o número do artigo depois do nome da lei. O [FONTE:id] serve para conferência e NÃO deve duplicar a citação. Se uma autoridade não existir na BASE, não a cite.
3. JURISPRUDÊNCIA: se a base não tiver precedentes, NÃO cite nenhum. Escreva [INSERIR JURISPRUDÊNCIA VERIFICADA] onde caberia.
4. Itens da base com status 'conferir' só têm RESUMO: use o conteúdo como ideia, mas NÃO invente número de artigo, inciso ou parágrafo além do que está em 'referencia', e NÃO coloque o texto entre aspas como se fosse transcrição literal.
5. PEDIDOS: cada pedido deve estar ligado a ao menos um [FONTE:id] e aos fatos. Não peça nada sem base nos fatos.
6. VALORES de indenização: não arbitre número; escreva [PREENCHER: valor do dano moral a critério do advogado].
7. Siga a estrutura do MODELO fornecido, completando-o. Texto formal, sem comentários seus fora da peça.
8. A seção 'Da alegação de tráfego aéreo' deve rebater a tese da ré apoiando-se no item da base [FONTE:cba_256_p3] (eventos precisam ser supervenientes, imprevisíveis e inevitáveis; a ré deve comprovar; impossibilidade de adotar medidas adequadas) SOMENTE com base nos itens da base e nos fatos."""


def _formatar_base(base):
    linhas = []
    for item in base:
        linhas.append(f"[id: {item['id']}] ({item['status']}) {item['referencia']}\n   {item['texto']}")
    return "\n".join(linhas)


def gerar_peca(modelo_texto, dados_caso, fatos, base, anonimizador=None):
    bloco = f"DADOS DO CASO:\n{dados_caso}\n\nFATOS (já revisados pelo advogado):\n{fatos}"
    mapa = {}
    if anonimizador is not None:
        bloco, mapa = anonimizador.mascarar(bloco)
    user = f"MODELO DA PEÇA:\n{modelo_texto}\n\n{bloco}\n\nBASE LEGAL (única fonte permitida):\n{_formatar_base(base)}"
    bruto = _chamar_ia(REGRAS_PECA, user, 8000)
    if anonimizador is not None:
        bruto = anonimizador.restaurar(bruto, mapa)
    return bruto


# ---------------------------------------------------------------- sugestão de jurisprudência
REGRAS_SUGERIR_JURIS = """Você ajuda um advogado brasileiro a LOCALIZAR jurisprudência e normas úteis ao caso.
Você NÃO redige a peça agora: apenas sugere fontes para o advogado CONFERIR na origem.
REGRAS ABSOLUTAS:
- NUNCA invente número de processo, súmula, tema ou data. Se não tiver certeza da identificação, escreva "identificação a conferir" e indique os termos exatos de busca.
- Traga no máximo 6 sugestões, nesta ordem, quando fizer sentido ao caso: (1) jurisprudência do tribunal local (TJ) se o Estado for identificável; (2) STJ (precedentes, recursos repetitivos); (3) súmulas (STJ/STF); (4) enunciados/normativos do CNJ e resoluções de agências reguladoras (ex.: ANAC).
- Baseie-se apenas nos fatos e nos temas da base legal fornecida.
Para cada sugestão, use EXATAMENTE este formato:
[TRIBUNAL] — [tipo: acórdão / súmula / enunciado]
Tese: (o que a fonte decide)
Identificação: (número e ano, se tiver certeza; caso contrário, "a conferir")
Busca: (termos exatos para pesquisar no site oficial)
Fonte: (site oficial provável, ex.: https://scon.stj.jus.br)
""".strip()


def sugerir_jurisprudencia(dados_caso, fatos, base, anonimizador=None):
    entrada = (
        f"DADOS DO CASO:\n{dados_caso}\n\nFATOS:\n{fatos}\n\n"
        f"TEMAS/LEIS JÁ NA BASE:\n{_formatar_base(base)}"
    )
    mapa = {}
    if anonimizador is not None:
        entrada, mapa = anonimizador.mascarar(entrada)
    saida = _chamar_ia(REGRAS_SUGERIR_JURIS, entrada, 4000)
    if anonimizador is not None:
        saida = anonimizador.restaurar(saida, mapa)
    return (saida or "").strip()


# ---------------------------------------------------------------- chat com a IA
REGRAS_CHAT = """Você é um assistente jurídico do escritório, conversando em português com o advogado.
Objetivo: responder dúvidas e ajudar a PENSAR e PESQUISAR sobre o processo e o Direito brasileiro.
REGRAS ABSOLUTAS:
- Responda de forma clara, objetiva e prática, como um assistente jurídico experiente.
- NUNCA invente número de processo, de lei, de súmula, de tema, nem dados do caso. Se não tiver
  certeza da identificação, diga que é preciso CONFERIR na fonte oficial e indique onde buscar.
- Trate como referência apenas o que estiver no CONTEXTO DO CASO fornecido; se faltar dado, peça.
- Você é apoio: o advogado revisa e decide tudo antes de usar.
- Responda apenas à última mensagem do advogado, mantendo o fio da conversa.""".strip()


def responder_chat(historico, contexto="", anonimizador=None):
    """Responde a uma conversa. ``historico`` = lista de {'role','content'}."""
    partes = []
    if contexto and contexto.strip():
        partes.append("CONTEXTO DO CASO (use só como referência; não invente dados):\n" + contexto.strip())
    conversa = []
    for m in historico or []:
        quem = "Advogado" if m.get("role") == "user" else "Assistente"
        conversa.append(f"{quem}: {m.get('content', '')}")
    partes.append("CONVERSA ATÉ AGORA:\n" + "\n\n".join(conversa))
    partes.append("Responda agora à última mensagem do Advogado.")
    entrada = "\n\n".join(partes)
    mapa = {}
    if anonimizador is not None:
        entrada, mapa = anonimizador.mascarar(entrada)
    saida = _chamar_ia(REGRAS_CHAT, entrada, 4000)
    if anonimizador is not None:
        saida = anonimizador.restaurar(saida, mapa)
    return (saida or "").strip()


# ---------------------------------------------------------------- anexo / triagem de documento
def extrair_texto_anexo(nome, dados):
    """Extrai o texto de um anexo digital (.txt/.md/.docx/.pdf) localmente (LGPD)."""
    ext = Path(nome or "").suffix.lower()
    if ext in (".txt", ".md", ".csv"):
        for enc in ("utf-8", "latin-1"):
            try:
                return dados.decode(enc)
            except UnicodeDecodeError:
                continue
        return dados.decode("utf-8", "replace")
    if ext == ".docx":
        import io

        from docx import Document

        doc = Document(io.BytesIO(dados))
        partes = [p.text for p in doc.paragraphs]
        for tabela in doc.tables:
            for linha in tabela.rows:
                partes.append(" | ".join(c.text for c in linha.cells))
        return "\n".join(partes)
    if ext == ".pdf":
        try:
            import io

            from pypdf import PdfReader
        except ImportError as e:
            raise RuntimeError("Para ler PDF instale: pip install pypdf") from e
        leitor = PdfReader(io.BytesIO(dados))
        return "\n".join((pag.extract_text() or "") for pag in leitor.pages)
    raise RuntimeError(f"Formato não suportado: '{ext}'. Use PDF, DOCX ou TXT.")


REGRAS_ANALISE = """Você faz TRIAGEM jurídica para um escritório de advocacia brasileiro.
Tarefa: ler o DOCUMENTO ANEXADO e produzir um RELATÓRIO DE TRIAGEM para o advogado decidir a peça cabível.
REGRAS ABSOLUTAS:
- Trabalhe SOMENTE com o texto do documento e com os dados do caso fornecidos. NÃO invente fatos, datas, valores, cláusulas ou números.
- NÃO afirme ilegalidade como certeza: aponte como POSSÍVEL irregularidade e cite o trecho exato do documento entre aspas.
- NÃO invente número de lei, artigo, súmula ou julgado. Sugira o que pesquisar e marque "conferir".
- Se um dispositivo já existir na BASE LEGAL fornecida, cite-o marcando [FONTE:id].
Responda em Markdown, exatamente com estas seções:
## 1. Tipo de documento
## 2. Peça cabível
(uma ou mais peças possíveis, com justificativa curta; se depender de prazo, escreva "verificar prazo")
## 3. Possíveis irregularidades / erros
(para cada item: descrição + trecho do documento entre aspas + por que seria irregular)
## 4. Fundamentação a conferir
(dispositivos/normas/jurisprudência a conferir, com termos de busca)
## 5. Próximos passos
(prazos a verificar, documentos a reunir, providências)
""".strip()


def analisar_documento(texto_anexo, dados_caso, base, anonimizador=None):
    entrada = (
        f"DADOS DO CASO:\n{dados_caso}\n\n"
        f"BASE LEGAL (pode citar marcando [FONTE:id]):\n{_formatar_base(base)}\n\n"
        f"DOCUMENTO ANEXADO:\n{texto_anexo}"
    )
    mapa = {}
    if anonimizador is not None:
        entrada, mapa = anonimizador.mascarar(entrada)
    saida = _chamar_ia(REGRAS_ANALISE, entrada, 6000)
    if anonimizador is not None:
        saida = anonimizador.restaurar(saida, mapa)
    return (saida or "").strip()


# ---------------------------------------------------------------- transcrição (voz)
def transcrever_audio(audio_bytes, mime_type="audio/wav"):
    """Transcreve fielmente o áudio do cliente (Fase 2 - microfone)."""
    if PROVEDOR == "anthropic":
        raise RuntimeError("A transcrição por voz está disponível apenas com o provedor Gemini.")
    from google.genai import types

    cliente = _cliente_gemini()
    parte = types.Part.from_bytes(data=audio_bytes, mime_type=mime_type)
    prompt = (
        "Transcreva fielmente, em português do Brasil, o áudio a seguir, que contém a narrativa "
        "de um cliente contando o caso a um advogado. Não resuma, não interprete, não acrescente "
        "nada: devolva apenas a transcrição literal do que foi dito."
    )
    resp = cliente.models.generate_content(model=MODELO, contents=[prompt, parte])
    return (resp.text or "").strip()


# ---------------------------------------------------------------- conferência (anti-invenção)
def _norm(s):
    return re.sub(r"\s+", " ", (s or "").lower()).strip()


def _so_digitos(s):
    return re.sub(r"\D", "", s or "")


_PADROES_SENSIVEIS = {
    "data": re.compile(r"\b\d{1,2}/\d{1,2}/\d{2,4}\b"),
    "valor": re.compile(r"R\$\s*\d{1,3}(?:\.\d{3})*(?:,\d{2})?"),
    "hora": re.compile(r"\b\d{1,2}(?::|h)\d{2}\b"),
}


def extrair_dados_sensiveis(texto):
    return {
        chave: sorted({m.group(0) for m in padrao.finditer(texto or "")})
        for chave, padrao in _PADROES_SENSIVEIS.items()
    }


def _dado_presente(token, fonte_texto):
    digitos = _so_digitos(token)
    return bool(digitos) and digitos in _so_digitos(fonte_texto)


def _verificar_dados_inventados(peca, fonte_fatos):
    alertas = []
    fonte = fonte_fatos or ""
    if not fonte.strip():
        return alertas
    na_peca = extrair_dados_sensiveis(peca)
    for categoria in ("data", "valor", "hora"):
        for dado in na_peca[categoria]:
            if not _dado_presente(dado, fonte):
                alertas.append((
                    "aviso",
                    f"{categoria.title()} na peça que NÃO aparece nos fatos fornecidos: '{dado}'. "
                    "Verifique se não foi inventado.",
                ))
    return alertas


def conferir_citacoes(texto, base, fonte_fatos=""):
    """
    Confere a peça contra a base legal e contra os fatos.
    Devolve dict: {texto, alertas: [(nivel, msg)], usadas: [itens]}.
    niveis: 'erro' (grave), 'aviso' (revisar), 'ok' (informativo).
    """
    por_id = {i["id"]: i for i in base}
    usadas, alertas = [], []

    def _render_citacao(item, antes):
        """Insere só a parte da referência que ainda não aparece no texto (evita repetir o artigo)."""
        ref = item["referencia"]
        if "–" in ref:
            art, lei = (x.strip() for x in ref.split("–", 1))
        else:
            art, lei = ref.strip(), ""
        nums_ref = set(re.findall(r"art[s]?\.?\s*(\d+)", art)) | set(re.findall(r"\be\s+(\d+)", art))
        nums_antes = set(re.findall(r"art[s]?\.?\s*(\d+)", antes))
        tem_art = bool(nums_ref & nums_antes)
        dig_lei = _so_digitos(lei)
        tem_lei = bool(dig_lei) and dig_lei in _so_digitos(antes)
        if tem_art and (tem_lei or not lei):
            return ""
        if tem_art and lei:
            return f" ({lei})"
        if tem_lei and art:
            return f" ({art})"
        return f" ({ref})"

    partes, pos = [], 0
    for m in re.finditer(r"\[FONTE:\s*([^\]]+)\]", texto):
        partes.append(texto[pos:m.start()])
        id_ = m.group(1).strip()
        antes = texto[max(0, m.start() - 200):m.start()]
        if id_ not in por_id:
            alertas.append(("erro", f"Fonte INEXISTENTE na base: '{id_}'. A IA pode ter inventado — remova ou cadastre a fonte."))
            _LOG.warning("Citação inexistente: %s", id_)
            partes.append(f"[FONTE INEXISTENTE: {id_}]")
        else:
            item = por_id[id_]
            if item not in usadas:
                usadas.append(item)
            partes.append(_render_citacao(item, antes))
        pos = m.end()
    partes.append(texto[pos:])

    final = "".join(partes)
    final = re.sub(r"\(\s*\)", "", final)
    final = re.sub(r"[ \t]{2,}", " ", final)
    final = re.sub(r"\s+([.,;:])", r"\1", final)
    final = re.sub(r"(\([^()]*\))(?:[\s,;.]+)\1", r"\1", final)

    base_ref = _norm(" ".join(i["referencia"] for i in base))

    # jurisprudência / súmulas / temas / leis fora da base
    padrao = r"(S[úu]mula(?: Vinculante)?\s*(?:n[ºo.]*\s*)?\d+|\b(?:A?REsp|ARE|RE|HC|AgInt|AgRg)\s*(?:n[ºo.]*\s*)?[\d\.\-]{3,}|Tema\s+\d+|Lei\s*(?:n[ºo.]*\s*)?[\d\.]+/\d+)"
    for m in sorted(set(re.findall(padrao, final, flags=re.I))):
        if _so_digitos(m) not in _so_digitos(base_ref):
            alertas.append(("aviso", f"Citação fora da base: '{m}' — confira na fonte oficial ou remova."))

    # artigos fora da base
    permitidos = set(re.findall(r"art[s]?\.?\s*(\d+)", base_ref)) | set(re.findall(r"\be\s+(\d+)", base_ref))
    for n in sorted(set(re.findall(r"\bart[s]?\.?\s*(\d+)", final, flags=re.I)), key=int):
        if n not in permitidos:
            alertas.append(("aviso", f"'art. {n}' não consta na base — confira na lei ou adicione à base."))

    # status das fontes efetivamente usadas
    for i in usadas:
        if i.get("status") != "verificado":
            alertas.append(("aviso", f"Fonte '{i['referencia']}' está com status '{i.get('status')}': confira o texto oficial."))

    # dados potencialmente inventados
    alertas.extend(_verificar_dados_inventados(final, fonte_fatos))

    # campos pendentes
    if "[PREENCHER" in final or "[INSERIR" in final:
        alertas.append(("aviso", "A peça tem campos [PREENCHER]/[INSERIR] pendentes."))

    # marcadores de anonimização que por algum motivo voltaram sem restaurar
    restantes = sorted(set(Anonimizador._RX_MARCADOR.findall(final)))
    if restantes:
        alertas.append(("erro", f"Marcador(es) de anonimização não restaurado(s): {', '.join(restantes)}. Confira esses trechos antes de protocolar."))

    if not alertas:
        alertas.append(("ok", "Nenhum alerta automático. Ainda assim, revise tudo antes de protocolar."))

    return {"texto": final, "alertas": alertas, "usadas": usadas}


# ---------------------------------------------------------------- exportação
_RE_TITULO = re.compile(
    r"^\s*(?:[IVX]+\s*[-–)]|[A-ZÇÃÕÁÉÍÓÚÂÊÔÀ][A-ZÇÃÕÁÉÍÓÚÂÊÔÀ\s]{8,}:?)\s*$"
)


def _eh_titulo(linha):
    if linha.lstrip().upper().startswith("AO JUÍZO"):
        return True
    return bool(_RE_TITULO.match(linha))


def salvar_docx(texto, caminho):
    """Exporta em .docx seguindo a formatação ABNT (NBR 14724)."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    secao = doc.sections[0]
    secao.page_width = Cm(21.0)          # A4
    secao.page_height = Cm(29.7)
    secao.top_margin = Cm(3.0)           # ABNT: margens 3/3/2/2 cm
    secao.left_margin = Cm(3.0)
    secao.right_margin = Cm(2.0)
    secao.bottom_margin = Cm(2.0)

    estilo = doc.styles["Normal"]
    estilo.font.name = "Times New Roman"
    estilo.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    estilo.font.size = Pt(12)
    estilo.paragraph_format.line_spacing = 1.5
    estilo.paragraph_format.space_after = Pt(6)

    for linha in (texto or "").split("\n"):
        conteudo = linha.rstrip()
        if not conteudo.strip():
            continue
        p = doc.add_paragraph()
        run = p.add_run(conteudo)
        if _eh_titulo(conteudo):
            run.bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(1.25)

    doc.save(str(caminho))
    _LOG.info("DOCX salvo: %s", caminho)
    _backup_drive(caminho, "pecas")
    return caminho


def _texto_latin1(texto):
    trocas = {"\u2013": "-", "\u2014": "-", "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"', "\u2026": "..."}
    for de, para in trocas.items():
        texto = texto.replace(de, para)
    return texto.encode("latin-1", "replace").decode("latin-1")


def salvar_pdf(texto, caminho):
    """Exporta em PDF seguindo a formatação ABNT (NBR 14724)."""
    try:
        from fpdf import FPDF
    except ImportError as e:
        raise RuntimeError("Para exportar PDF instale: pip install fpdf2") from e

    pdf = FPDF(format="A4", unit="mm")
    pdf.set_margins(30, 30, 20)          # ABNT: esq/topo 3 cm, dir 2 cm
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()
    pdf.set_font("Times", size=12)
    largura = pdf.epw
    recuo = 12.5                          # primeira linha: 1,25 cm
    altura = 7                            # entrelinha ~1,5
    for linha in (texto or "").split("\n"):
        conteudo = linha.rstrip()
        if not conteudo.strip():
            pdf.ln(altura)
            continue
        seguro = _texto_latin1(conteudo)
        if _eh_titulo(conteudo):
            pdf.set_font("Times", style="B", size=12)
            pdf.multi_cell(largura, altura, seguro, align="L")
            pdf.set_font("Times", size=12)
        else:
            pdf.set_x(pdf.l_margin + recuo)
            pdf.multi_cell(largura - recuo, altura, seguro, align="J")

    pdf.output(str(caminho))
    _LOG.info("PDF salvo: %s", caminho)
    _backup_drive(caminho, "pecas")
    return caminho


# ---------------------------------------------------------------- histórico de casos
def _dono_seguro(dono):
    return re.sub(r"[^A-Za-z0-9_\-]", "_", (dono or "").strip()).strip("_")


def _pasta_casos_local(dono):
    seguro = _dono_seguro(dono)
    if not seguro:
        return PASTA_HISTORICO
    pasta = PASTA_HISTORICO / seguro
    pasta.mkdir(exist_ok=True)
    return pasta


def _subpasta_casos(dono):
    seguro = _dono_seguro(dono)
    return f"casos/{seguro}" if seguro else "casos"


def salvar_caso(nome, dados, dono=""):
    seguro = re.sub(r"[^A-Za-z0-9_\-]", "_", nome or "").strip("_") or datetime.now().strftime("%Y%m%d_%H%M%S")
    caminho = _pasta_casos_local(dono) / f"{seguro}.json"
    dados = dict(dados)
    dados["_salvo_em"] = datetime.now().isoformat(timespec="seconds")
    dados["_dono"] = (dono or "").strip()
    caminho.write_text(json.dumps(dados, ensure_ascii=False, indent=2), encoding="utf-8")
    _backup_drive(caminho, _subpasta_casos(dono))
    return caminho


def listar_casos(dono=""):
    nomes = set()
    pasta = _pasta_casos_local(dono)
    if pasta.exists():
        nomes.update(p.stem for p in pasta.glob("*.json"))
    try:
        import nuvem

        if nuvem.disponivel():
            for nome_arq, tipo in nuvem.listar(_subpasta_casos(dono)):
                if nome_arq and nome_arq.endswith(".json") and (tipo in (None, "file")):
                    nomes.add(nome_arq[:-5])
    except Exception:
        _LOG.warning("Listar casos na nuvem falhou (dono=%s)", dono, exc_info=True)
    return sorted(nomes, reverse=True)


def carregar_caso(nome, dono=""):
    caminho = _pasta_casos_local(dono) / f"{nome}.json"
    if caminho.exists():
        with open(caminho, encoding="utf-8") as f:
            return json.load(f)
    try:
        import nuvem

        if nuvem.disponivel():
            bruto = nuvem.ler_texto(f"{_subpasta_casos(dono)}/{nome}.json")
            if bruto:
                caminho.write_text(bruto, encoding="utf-8")
                return json.loads(bruto)
    except Exception:
        _LOG.warning("Buscar caso na nuvem falhou (%s)", nome, exc_info=True)
    raise FileNotFoundError(f"Caso não encontrado: {nome}")


def apagar_caso(nome, dono=""):
    caminho = _pasta_casos_local(dono) / f"{nome}.json"
    if caminho.exists():
        caminho.unlink()
    _backup_remover(f"{nome}.json", _subpasta_casos(dono))


# ------------------------------------------------------------------ rascunho automático
def _pasta_rascunho(dono, criar=True):
    pasta = PASTA_RASCUNHOS / (_dono_seguro(dono) or "sem_dono")
    if criar:
        pasta.mkdir(parents=True, exist_ok=True)
    return pasta


def salvar_rascunho(dados, dono=""):
    """Guarda o trabalho em andamento (relato/fatos/peça) para nunca perder.

    É o que salva a página se o app fechar, o navegador travar ou a sessão cair.
    """
    caminho = _pasta_rascunho(dono) / "rascunho.json"
    dados = dict(dados)
    dados["_salvo_em"] = datetime.now().isoformat(timespec="seconds")
    dados["_dono"] = (dono or "").strip()
    caminho.write_text(json.dumps(dados, ensure_ascii=False, indent=2), encoding="utf-8")
    return caminho


def carregar_rascunho(dono=""):
    """Devolve o rascunho salvo (ou None)."""
    caminho = _pasta_rascunho(dono, criar=False) / "rascunho.json"
    if not caminho.exists():
        return None
    try:
        with open(caminho, encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        _LOG.warning("Rascunho ilegível (%s)", caminho, exc_info=True)
        return None


def apagar_rascunho(dono=""):
    """Apaga o rascunho (quando o trabalho foi salvo como caso)."""
    caminho = _pasta_rascunho(dono, criar=False) / "rascunho.json"
    try:
        if caminho.exists():
            caminho.unlink()
        return True
    except Exception:
        return False


def _texto_para_ia(texto):
    """Tira as marcações visuais e deixa o texto pronto para a IA.

    O lawyer marca o texto (negrito, sublinhado, destaque) só para highlight;
    para a IA, negrito e itálico em markdown bastam.
    """
    t = texto or ""
    t = re.sub(r"==(.+?)==", r"**\1**", t, flags=re.S)
    t = re.sub(r"__(.+?)__", r"\1", t, flags=re.S)
    return t
